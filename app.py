import os
import base64
import cv2
import numpy as np
import openpyxl
from datetime import datetime
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file, send_from_directory
from flask_sqlalchemy import SQLAlchemy
import face_recognition

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "visionpass_attendance_secret_2026")
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///users.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)

# User model with Role-Based Access Control (RBAC)
class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)
    role = db.Column(db.String(20), default="student", nullable=False)  # 'student', 'teacher', 'admin'

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

# Paths and global state
attendance_file = "attendance.xlsx"
images_dir = 'known_faces'
os.makedirs(images_dir, exist_ok=True)

known_face_encodings = []
known_face_names = []

def load_known_faces():
    """Load and encode all reference faces in the known_faces directory."""
    global known_face_encodings, known_face_names
    known_face_encodings = []
    known_face_names = []

    if not os.path.exists(images_dir):
        return

    valid_extensions = ('.jpg', '.jpeg', '.png', '.webp')
    for image_name in os.listdir(images_dir):
        if not image_name.lower().endswith(valid_extensions):
            continue

        image_path = os.path.join(images_dir, image_name)
        try:
            image = face_recognition.load_image_file(image_path)
            encodings = face_recognition.face_encodings(image)
            if len(encodings) > 0:
                known_face_encodings.append(encodings[0])
                name = os.path.splitext(image_name)[0]
                known_face_names.append(name)
                print(f"[Loaded Face] Identity: {name}")
            else:
                print(f"[Warning] No detectable face found in {image_name}")
        except Exception as e:
            print(f"[Error] Failed to process {image_name}: {e}")

load_known_faces()

# Excel Attendance Helpers
def init_attendance_file():
    if not os.path.exists(attendance_file):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(["Name", "Date", "Time"])
        wb.save(attendance_file)

init_attendance_file()

def get_user_attendance(username):
    """Retrieve attendance records for a specific user from attendance.xlsx."""
    records = []
    if not os.path.exists(attendance_file):
        return records

    try:
        wb = openpyxl.load_workbook(attendance_file, data_only=True)
        ws = wb.active
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row and len(row) >= 3 and row[0]:
                record_name = str(row[0]).strip()
                if record_name.lower() == username.lower():
                    records.append({
                        "row_index": idx,
                        "name": record_name,
                        "date": str(row[1]) if row[1] else "",
                        "time": str(row[2]) if row[2] else ""
                    })
    except Exception as e:
        print(f"[Error] Reading attendance.xlsx: {e}")

    return list(reversed(records))

def get_all_attendance():
    """Retrieve all attendance records with Excel row index."""
    records = []
    if not os.path.exists(attendance_file):
        return records

    try:
        wb = openpyxl.load_workbook(attendance_file, data_only=True)
        ws = wb.active
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row and len(row) >= 3 and row[0]:
                records.append({
                    "row_index": idx,
                    "name": str(row[0]).strip(),
                    "date": str(row[1]) if row[1] else "",
                    "time": str(row[2]) if row[2] else ""
                })
    except Exception as e:
        print(f"[Error] Reading all attendance: {e}")

    return list(reversed(records))

def record_attendance(name):
    """Append a verified attendance record into attendance.xlsx."""
    now = datetime.now()
    date_str = now.strftime("%Y-%m-%d")
    time_str = now.strftime("%H:%M:%S")

    wb = openpyxl.load_workbook(attendance_file)
    ws = wb.active
    ws.append([name, date_str, time_str])
    wb.save(attendance_file)
    return date_str, time_str

def delete_attendance_row(row_index):
    """Delete a specific row from attendance.xlsx by 1-indexed row number."""
    if not os.path.exists(attendance_file):
        return False
    try:
        wb = openpyxl.load_workbook(attendance_file)
        ws = wb.active
        ws.delete_rows(row_index, 1)
        wb.save(attendance_file)
        return True
    except Exception as e:
        print(f"[Error] Deleting attendance row {row_index}: {e}")
        return False

def clear_all_attendance():
    """Clear all attendance logs while keeping the headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(["Name", "Date", "Time"])
    wb.save(attendance_file)

def process_face_image(rgb_image, logged_in_username):
    """Core face matching logic across loaded known faces."""
    face_locations = face_recognition.face_locations(rgb_image)
    if not face_locations:
        return False, "No face detected in the frame. Please look directly at the camera with clear lighting."

    face_encodings = face_recognition.face_encodings(rgb_image, face_locations)
    if not face_encodings:
        return False, "Could not extract facial biometric features. Please try again."

    if not known_face_encodings:
        return False, "No enrolled faces registered in the system database yet."

    for face_encoding in face_encodings:
        matches = face_recognition.compare_faces(known_face_encodings, face_encoding, tolerance=0.52)
        face_distances = face_recognition.face_distance(known_face_encodings, face_encoding)
        best_match_index = int(np.argmin(face_distances))

        if matches[best_match_index]:
            recognized_name = known_face_names[best_match_index]

            if recognized_name.lower() == logged_in_username.lower():
                date_str, time_str = record_attendance(recognized_name)
                return True, f"Attendance successfully marked for {recognized_name} on {date_str} at {time_str}!"
            else:
                return False, f"Face detected as '{recognized_name}', which does NOT match the logged-in user '{logged_in_username}'."

    return False, "Face not recognized. Please ensure your reference face photo is enrolled."

# Access Control Decorators
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "username" not in session:
            flash("Please sign in to access this page.", "info")
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated

def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "username" not in session:
            flash("Please sign in with administrator credentials.", "info")
            return redirect(url_for("login"))
        if session.get("role") != "admin":
            flash("Access denied. Master Administrator privileges required.", "error")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return decorated

def teacher_or_admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "username" not in session:
            flash("Please sign in to access the management portal.", "info")
            return redirect(url_for("login"))
        if session.get("role") not in ["teacher", "admin"]:
            flash("Access denied. Teacher or Administrator privileges required.", "error")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return decorated

# Context Processor for Global Template Data
@app.context_processor
def inject_user():
    return {
        "current_user": session.get("username"),
        "current_role": session.get("role")
    }

# ==================== PUBLIC ROUTES ====================

@app.route("/favicon.ico")
def favicon():
    return send_from_directory(os.path.join(app.root_path, 'static'), 'favicon.ico', mimetype='image/vnd.microsoft.icon')

@app.route("/")
def home():
    if "username" in session:
        role = session.get("role", "student")
        if role == "admin":
            return redirect(url_for("admin_dashboard"))
        elif role == "teacher":
            return redirect(url_for("teacher_dashboard"))
        return redirect(url_for("dashboard"))
    return render_template("index.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()

        user = User.query.filter_by(username=username).first()

        if user and user.check_password(password):
            session["username"] = user.username
            session["role"] = getattr(user, "role", "student")
            flash(f"Welcome back, {user.username}! ({session['role'].capitalize()} access)", "success")
            
            if session["role"] == "admin":
                return redirect(url_for("admin_dashboard"))
            elif session["role"] == "teacher":
                return redirect(url_for("teacher_dashboard"))
            else:
                return redirect(url_for("dashboard"))
        else:
            flash("Invalid credentials. Please verify your username and password.", "error")
            return redirect(url_for("login"))

    return render_template("login.html")

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()

        if not username or not password:
            flash("Username and password are required.", "error")
            return redirect(url_for("register"))

        if User.query.filter_by(username=username).first():
            flash("Username already exists. Please choose a different one.", "error")
            return redirect(url_for("register"))

        # Strictly enforce Face Biometric Capture (required)
        file = request.files.get("face_image")
        webcam_data = request.form.get("webcam_face_data")

        has_webcam = bool(webcam_data and "base64," in webcam_data)
        has_file = bool(file and file.filename != "")

        if not has_webcam and not has_file:
            flash("Face Biometric Capture is required. Please capture a camera snapshot or upload a face photo.", "error")
            return redirect(url_for("register"))

        target_path = None
        try:
            if has_webcam:
                base64_str = webcam_data.split("base64,")[1]
                img_bytes = base64.b64decode(base64_str)
                target_filename = f"{secure_filename(username)}.jpg"
                target_path = os.path.join(images_dir, target_filename)
                with open(target_path, "wb") as f:
                    f.write(img_bytes)

            elif has_file:
                ext = os.path.splitext(file.filename)[1].lower()
                if ext not in ['.jpg', '.jpeg', '.png', '.webp']:
                    flash("Invalid image format. Please upload JPG, PNG, or WEBP.", "error")
                    return redirect(url_for("register"))
                target_filename = f"{secure_filename(username)}{ext}"
                target_path = os.path.join(images_dir, target_filename)
                file.save(target_path)

            # Validate face in image
            loaded_img = face_recognition.load_image_file(target_path)
            encs = face_recognition.face_encodings(loaded_img)

            if len(encs) == 0:
                if target_path and os.path.exists(target_path):
                    os.remove(target_path)
                flash("Face verification failed: No detectable face found. A clear face photo is required to enroll.", "error")
                return redirect(url_for("register"))

        except Exception as e:
            if target_path and os.path.exists(target_path):
                try:
                    os.remove(target_path)
                except Exception:
                    pass
            flash(f"Biometric processing error: {str(e)}", "error")
            return redirect(url_for("register"))

        # Create user account ONLY after face biometric is successfully verified
        new_user = User(username=username, role="student")
        new_user.set_password(password)
        db.session.add(new_user)
        db.session.commit()

        # Update in-memory face encodings
        load_known_faces()

        flash(f"Registration successful! Biometric profile enrolled for {username}. Please sign in.", "success")
        return redirect(url_for("login"))

    return render_template("register.html")

@app.route("/logout")
def logout():
    session.clear()
    flash("You have been signed out.", "info")
    return redirect(url_for("login"))

# ==================== STUDENT / ATTENDEE PORTAL ====================

@app.route("/dashboard")
@login_required
def dashboard():
    username = session["username"]
    records = get_user_attendance(username)
    has_face = username.lower() in [n.lower() for n in known_face_names]

    return render_template("dashboard.html", username=username, attendance_records=records, has_face=has_face)

@app.route("/mark_attendance", methods=["GET", "POST"])
@login_required
def mark_attendance():
    username = session["username"]

    if request.method == "POST":
        image_data = request.form.get("image_data")
        mode = request.form.get("mode")

        if image_data and "base64," in image_data:
            try:
                base64_str = image_data.split("base64,")[1]
                img_bytes = base64.b64decode(base64_str)
                nparr = np.frombuffer(img_bytes, np.uint8)
                frame = cv2.imdecode(nparr, cv2.IMREAD_COLOR)

                if frame is None:
                    flash("Failed to decode webcam snapshot.", "error")
                    return redirect(url_for("mark_attendance"))

                rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                success, message = process_face_image(rgb_frame, username)

                if success:
                    flash(message, "success")
                    role = session.get("role", "student")
                    if role == "admin":
                        return redirect(url_for("admin_dashboard"))
                    elif role == "teacher":
                        return redirect(url_for("teacher_dashboard"))
                    return redirect(url_for("dashboard"))
                else:
                    flash(message, "error")
                    return redirect(url_for("mark_attendance"))
            except Exception as e:
                flash(f"Processing error: {str(e)}", "error")
                return redirect(url_for("mark_attendance"))

        elif mode == "server":
            video_capture = cv2.VideoCapture(0)
            if not video_capture.isOpened():
                flash("Error: Could not access server webcam hardware.", "error")
                return redirect(url_for("mark_attendance"))

            try:
                for _ in range(5):
                    video_capture.read()

                ret, frame = video_capture.read()
                video_capture.release()

                if not ret or frame is None:
                    flash("Error: Could not capture clear frame from server camera.", "error")
                    return redirect(url_for("mark_attendance"))

                small_frame = cv2.resize(frame, (0, 0), fx=0.5, fy=0.5)
                rgb_small_frame = cv2.cvtColor(small_frame, cv2.COLOR_BGR2RGB)

                success, message = process_face_image(rgb_small_frame, username)
                if success:
                    flash(message, "success")
                    return redirect(url_for("dashboard"))
                else:
                    flash(message, "error")
                    return redirect(url_for("mark_attendance"))
            except Exception as e:
                if video_capture.isOpened():
                    video_capture.release()
                flash(f"Camera capture error: {str(e)}", "error")
                return redirect(url_for("mark_attendance"))

        else:
            flash("No camera feed data received.", "error")
            return redirect(url_for("mark_attendance"))

    return render_template("mark_attendance.html", username=username)

@app.route("/download_attendance")
@login_required
def download_attendance():
    if os.path.exists(attendance_file):
        return send_file(attendance_file, as_attachment=True, download_name="attendance.xlsx")
    flash("Attendance log file not found.", "error")
    return redirect(url_for("dashboard"))

# ==================== TEACHER / MANAGER PORTAL ====================

@app.route("/teacher/dashboard")
@teacher_or_admin_required
def teacher_dashboard():
    all_logs = get_all_attendance()
    students = User.query.filter_by(role="student").all()
    today_str = datetime.now().strftime("%Y-%m-%d")

    # Calculate statistics
    today_logs = [log for log in all_logs if log["date"] == today_str]
    unique_attendees_today = len(set(log["name"].lower() for log in today_logs))
    total_students = len(students)
    attendance_rate = f"{(unique_attendees_today / total_students * 100):.1f}%" if total_students > 0 else "0%"

    return render_template(
        "teacher_dashboard.html",
        username=session["username"],
        role=session["role"],
        all_logs=all_logs,
        students=students,
        today_count=len(today_logs),
        unique_attendees_today=unique_attendees_today,
        total_students=total_students,
        attendance_rate=attendance_rate
    )

@app.route("/teacher/manual_attendance", methods=["POST"])
@teacher_or_admin_required
def teacher_manual_attendance():
    student_name = request.form.get("student_name", "").strip()
    if not student_name:
        flash("Student name is required for manual attendance.", "error")
        return redirect(url_for("teacher_dashboard"))

    date_str, time_str = record_attendance(student_name)
    flash(f"Manual attendance marked for {student_name} on {date_str} at {time_str} by {session['username']}.", "success")
    return redirect(url_for("teacher_dashboard"))

# ==================== MASTER ADMIN PORTAL ====================

@app.route("/admin/dashboard")
@admin_required
def admin_dashboard():
    users = User.query.all()
    all_logs = get_all_attendance()

    # Get enrolled face files
    face_files = []
    if os.path.exists(images_dir):
        face_files = [f for f in os.listdir(images_dir) if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp'))]

    role_counts = {
        "admin": sum(1 for u in users if getattr(u, 'role', 'student') == "admin"),
        "teacher": sum(1 for u in users if getattr(u, 'role', 'student') == "teacher"),
        "student": sum(1 for u in users if getattr(u, 'role', 'student') == "student"),
    }

    return render_template(
        "admin_dashboard.html",
        username=session["username"],
        users=users,
        all_logs=all_logs,
        face_files=face_files,
        role_counts=role_counts,
        total_encodings=len(known_face_encodings)
    )

@app.route("/admin/users/create", methods=["POST"])
@admin_required
def admin_create_user():
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "").strip()
    role = request.form.get("role", "student").strip()

    if not username or not password:
        flash("Username and password are required.", "error")
        return redirect(url_for("admin_dashboard"))

    if User.query.filter_by(username=username).first():
        flash("Username already exists.", "error")
        return redirect(url_for("admin_dashboard"))

    new_user = User(username=username, role=role)
    new_user.set_password(password)
    db.session.add(new_user)
    db.session.commit()

    flash(f"Account for {username} ({role.capitalize()}) created successfully.", "success")
    return redirect(url_for("admin_dashboard"))

@app.route("/admin/users/update_role/<int:user_id>", methods=["POST"])
@admin_required
def admin_update_role(user_id):
    user = User.query.get(user_id)
    if not user:
        flash("User not found.", "error")
        return redirect(url_for("admin_dashboard"))

    new_role = request.form.get("role")
    if new_role in ["student", "teacher", "admin"]:
        user.role = new_role
        db.session.commit()
        flash(f"Updated role for {user.username} to {new_role.capitalize()}.", "success")
    return redirect(url_for("admin_dashboard"))

@app.route("/admin/users/delete/<int:user_id>", methods=["POST"])
@admin_required
def admin_delete_user(user_id):
    user = User.query.get(user_id)
    if not user:
        flash("User not found.", "error")
        return redirect(url_for("admin_dashboard"))

    if user.username == session["username"]:
        flash("You cannot delete your own active administrator account!", "error")
        return redirect(url_for("admin_dashboard"))

    username = user.username
    db.session.delete(user)
    db.session.commit()

    # Also clean up associated face image if exists
    for ext in ['.jpg', '.jpeg', '.png', '.webp']:
        p = os.path.join(images_dir, f"{username}{ext}")
        if os.path.exists(p):
            try:
                os.remove(p)
            except Exception:
                pass
    load_known_faces()

    flash(f"User {username} and associated biometric records deleted.", "success")
    return redirect(url_for("admin_dashboard"))

@app.route("/admin/attendance/delete_row", methods=["POST"])
@admin_required
def admin_delete_attendance_row():
    try:
        row_idx = int(request.form.get("row_index"))
        if delete_attendance_row(row_idx):
            flash(f"Deleted attendance record at row {row_idx}.", "success")
        else:
            flash("Failed to delete record.", "error")
    except Exception as e:
        flash(f"Error: {str(e)}", "error")
    return redirect(url_for("admin_dashboard"))

@app.route("/admin/attendance/clear_all", methods=["POST"])
@admin_required
def admin_clear_all_attendance():
    clear_all_attendance()
    flash("All attendance records have been cleared from attendance.xlsx.", "success")
    return redirect(url_for("admin_dashboard"))

@app.route("/admin/faces/delete/<filename>", methods=["POST"])
@admin_required
def admin_delete_face(filename):
    file_path = os.path.join(images_dir, secure_filename(filename))
    if os.path.exists(file_path):
        os.remove(file_path)
        load_known_faces()
        flash(f"Removed biometric face file {filename}.", "success")
    else:
        flash("Face file not found.", "error")
    return redirect(url_for("admin_dashboard"))

@app.route("/face_image/<filename>")
@teacher_or_admin_required
def serve_face_image(filename):
    """Serve enrolled face thumbnail for admin/teacher view."""
    return send_from_directory(images_dir, secure_filename(filename))

# Database Initialization & Seeding
def init_db_and_seed():
    with app.app_context():
        db.create_all()
        # Automatic SQLite column migration for 'role'
        try:
            engine = db.engine
            with engine.connect() as conn:
                res = conn.exec_driver_sql("PRAGMA table_info(user)").fetchall()
                col_names = [r[1] for r in res]
                if "role" not in col_names:
                    conn.exec_driver_sql("ALTER TABLE user ADD COLUMN role VARCHAR(20) DEFAULT 'student'")
                    conn.commit()
                    print("[Database] Successfully added 'role' column to 'user' table.")
        except Exception as e:
            print(f"[Database Migration Notice] {e}")

        # Seed admin / admin123
        admin = User.query.filter_by(username="admin").first()
        if not admin:
            admin = User(username="admin", role="admin")
            admin.set_password("admin123")
            db.session.add(admin)
            db.session.commit()
            print("[Database] Seeded Master Administrator: 'admin' / 'admin123'")
        elif admin.role != "admin":
            admin.role = "admin"
            db.session.commit()

        # Seed teacher / teacher123
        teacher = User.query.filter_by(username="teacher").first()
        if not teacher:
            teacher = User(username="teacher", role="teacher")
            teacher.set_password("teacher123")
            db.session.add(teacher)
            db.session.commit()
            print("[Database] Seeded Teacher/Manager: 'teacher' / 'teacher123'")
        elif teacher.role != "teacher":
            teacher.role = "teacher"
            db.session.commit()

        # Seed student / student123
        student = User.query.filter_by(username="student").first()
        if not student:
            student = User(username="student", role="student")
            student.set_password("student123")
            db.session.add(student)
            db.session.commit()
            print("[Database] Seeded Student: 'student' / 'student123'")

        # Seed nihar
        nihar = User.query.filter_by(username="nihar").first()
        if not nihar:
            nihar = User(username="nihar", role="student")
            nihar.set_password("nihar123")
            db.session.add(nihar)
            db.session.commit()
            print("[Database] Seeded Student: 'nihar' / 'nihar123'")

init_db_and_seed()

if __name__ == "__main__":
    print("\n" + "="*50)
    print(" VISIONPASS MULTI-ROLE ATTENDANCE SYSTEM IS LIVE")
    print(" Local URL: http://127.0.0.1:5000")
    print(" Accounts:")
    print("   Admin:   admin / admin123")
    print("   Teacher: teacher / teacher123")
    print("   Student: student / student123")
    print("="*50 + "\n")
    app.run(host="0.0.0.0", port=5000, debug=True)